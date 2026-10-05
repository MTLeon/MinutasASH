from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from src.documents.anglo_minutes import AngloMinutesDocument
from src.models import MeetingItem, MeetingMetadata, MinuteAnalysis


def _template(path: Path) -> None:
    document = Document()
    header = document.add_table(rows=10, cols=6)
    header.cell(0, 0).text = "TEMA:"
    header.cell(1, 4).text = "N° MINUTA"
    header.cell(8, 0).text = "OBJETIVO DE LA REUNIÓN:"
    matrix = document.add_table(rows=20, cols=4)
    matrix.cell(0, 0).text = "ITEM N°"
    matrix.cell(0, 1).text = "DESCRIPCIÓN"
    matrix.cell(0, 2).text = "ACCIÓN / POR"
    matrix.cell(0, 3).text = "FECHA / CIERRE"
    document.add_table(rows=1, cols=3)
    document.save(path)


def test_anglo_provider_preserves_template_and_writes_reviewed_items() -> None:
    with TemporaryDirectory() as temporary:
        root = Path(temporary)
        template = root / "anglo.docx"
        output = root / "output.docx"
        _template(template)
        metadata = MeetingMetadata(
            minute_number="3165-MRE-PR-01",
            matter="Seguimiento Booster",
            project_code="3165",
        )
        analysis = MinuteAnalysis(
            objective="Cerrar acciones de ingeniería.",
            items=[
                MeetingItem(
                    category="compromiso",
                    description="Ivan Tapia enviará la revisión de ingeniería.",
                    responsible="Ivan Tapia",
                    due_date_iso="2026-10-02",
                    review_status="aprobado",
                ),
                MeetingItem(
                    category="informativo",
                    description="Comentario descartado.",
                    review_status="descartado",
                ),
            ],
        )
        AngloMinutesDocument().generate(
            analysis, metadata, output, {"anglo_template_path": template}
        )
        emitted = Document(output)
        text = "\n".join(
            cell.text for table in emitted.tables for row in table.rows for cell in row.cells
        )
        assert "3165-MRE-PR-01" in text
        assert "Ivan Tapia enviará la revisión de ingeniería." in text
        assert "Comentario descartado." not in text
