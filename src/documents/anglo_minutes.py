from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument

from src.anglo_sections import classify_anglo_section
from src.models import MeetingItem, MeetingMetadata, MinuteAnalysis
from src.review_quality import items_for_document


class AngloMinutesDocument:
    """Emite la minuta operacional de Anglo desde su formato DOCX oficial."""

    provider_id = "anglo_minutes_v1"
    display_name = "Minuta operacional Anglo American Sur"

    @staticmethod
    def _set(cell, value: str) -> None:
        cell.text = value

    @staticmethod
    def _date(item) -> str:
        return item.due_date_iso or item.due_date_text or "N.A."

    @staticmethod
    def _validate(
        document: DocxDocument, metadata: MeetingMetadata, analysis: MinuteAnalysis
    ) -> None:
        if len(document.tables) < 3:
            raise ValueError("La minuta Anglo generada no conserva sus tablas principales.")
        text = "\n".join(
            cell.text for table in document.tables for row in table.rows for cell in row.cells
        )
        if metadata.minute_number and metadata.minute_number not in text:
            raise ValueError("La minuta Anglo no contiene el número de minuta emitido.")
        active = items_for_document(analysis.items)
        if active and active[0].description not in text:
            raise ValueError("La minuta Anglo no contiene los puntos aprobados.")

    def generate(
        self,
        analysis: MinuteAnalysis,
        metadata: MeetingMetadata,
        output_path: Path,
        config: dict,
    ) -> Path:
        template_path = Path(str(config.get("anglo_template_path") or ""))
        if not template_path.is_file() or template_path.suffix.lower() != ".docx":
            raise FileNotFoundError("No está disponible la plantilla DOCX Anglo American Sur.")

        document = Document(str(template_path))
        if len(document.tables) < 2:
            raise ValueError("La plantilla Anglo no contiene la matriz de acuerdos requerida.")

        header = document.tables[0]
        if len(header.rows) >= 10:
            self._set(header.cell(0, 0), f"TEMA: {metadata.matter or 'Reunión de avance'}")
            self._set(header.cell(1, 4), metadata.minute_number or "Por confirmar")
            self._set(
                header.cell(8, 0),
                "OBJETIVO DE LA REUNIÓN: "
                + (analysis.objective or metadata.matter or "Por confirmar"),
            )

        matrix = document.tables[1]
        row_specs = {
            "HSE": [(2, "1.1"), (3, "1.2")],
            "Control y Programación": [(5, "2.1"), (6, "2.2"), (7, "2.3")],
            "Calidad": [(9, "3.1")],
            "Ingeniería": [(11, "4.1")],
            "Otros/General": [(13, "5.1"), (14, "5.2"), (15, "5.3"), (16, "5.4"), (17, "5.5")],
        }
        section_numbers = {
            "HSE": 1,
            "Control y Programación": 2,
            "Calidad": 3,
            "Ingeniería": 4,
            "Otros/General": 5,
        }
        for row_index, _code in (spec for rows in row_specs.values() for spec in rows):
            row = matrix.rows[row_index]
            for cell in row.cells:
                self._set(cell, "")
        active_items = items_for_document(analysis.items)
        grouped: dict[str, list[MeetingItem]] = {section: [] for section in row_specs}
        for item in active_items:
            grouped[classify_anglo_section(item)].append(item)
        for section, items in grouped.items():
            for index, item in enumerate(items):
                row, code = (
                    (matrix.rows[row_specs[section][index][0]], row_specs[section][index][1])
                    if index < len(row_specs[section])
                    else (matrix.add_row(), f"{section_numbers[section]}.{index + 1}")
                )
                self._set(row.cells[0], code)
                self._set(row.cells[1], item.description)
                self._set(row.cells[2], item.responsible or "Por confirmar")
                self._set(row.cells[3], self._date(item))

        if analysis.next_meeting and len(matrix.rows) >= 20:
            detail = " ".join(
                part
                for part in (
                    analysis.next_meeting.description,
                    analysis.next_meeting.date_text,
                    analysis.next_meeting.time_text,
                )
                if part
            )
            self._set(
                matrix.rows[19].cells[1], "Fecha Próxima Reunión: " + (detail or "Por confirmar")
            )

        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.core_properties.title = metadata.minute_number or "Minuta Anglo American Sur"
        document.core_properties.subject = metadata.matter or "Minuta de reunión"
        document.save(str(output_path))
        self._validate(Document(str(output_path)), metadata, analysis)
        return output_path
